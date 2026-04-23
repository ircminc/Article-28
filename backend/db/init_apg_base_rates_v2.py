"""Load APG base rates from the PMTAC 'Updated APG Fee Calculator' workbook.

Source: `Updated APG Fee Calculator - 04232026.xlsx`, sheet
`'Updated APG Base Rate'`. Wide-format layout:

    row 0: section title ('Freestanding Clinic and Ambulatory Surgery...')
    row 1: blank
    row 2: header row
        col 0 = 'Peer Group'
        col 1 = '**Base Rate Code'
        col 2 = '***Blend Rate Code'
        col 3 = 'Capital Rate Code'
        col 4 = 'Region'
        col 5 = compound key (ignored)
        col 6+ = effective-date columns (datetime cells OR strings like
                 '4/1/2022****' where asterisks are footnote markers)
    rows 3+: one row per (peer_group, region) combination. Rate cells may be
             empty for effective dates that predate the peer group's inception
             (e.g. the 'Episode' peer groups only start at 7/1/2011).
    end-of-data: first row whose column A is empty or starts with '*' (footnote).

Every non-empty (rate cell) becomes one `ApgBaseRate` row. We REPLACE every
row where `source='dtc'` — this sheet covers the full set of Freestanding
APG peer groups (Clinic*, Amb Surg, Renal, Academic Dental, Clinic MR/DD/TBI,
School-Based Health Center *, and all matching Episode rates). Hospital
base rates (`source='hospital'`) are left untouched.

The engine's `lookup_base_rate` already selects by exact peer-group match and
most-recent effective date <= DOS, so once the 2022-04-01 rates are loaded
here, every DOS >= 2022-04-01 will resolve to those rates automatically —
meeting both requirements of the user's update brief.
"""
from __future__ import annotations

import io
import logging
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Optional

import openpyxl
from sqlalchemy import delete
from sqlalchemy.ext.asyncio import AsyncSession

from backend.db.database import ApgBaseRate

log = logging.getLogger(__name__)


SHEET_NAME = "Updated APG Base Rate"
HEADER_ROW = 2          # 0-indexed: third row of the sheet
FIRST_DATE_COL = 6      # 0-indexed: first effective-date column
DATA_START_ROW = 3

# After the data ends the sheet has footnotes keyed by asterisks (e.g.
# "*For Clinic and ..."). Any column-A value that starts with '*' is NOT
# a peer group.
_FOOTNOTE_PREFIX = "*"


def _parse_date_header(cell) -> Optional[date]:
    """Return a date for a header cell, or None if it isn't a date.

    openpyxl returns datetime objects for date-formatted cells. Some cells
    in this workbook are string constants like '4/1/2022****' (the asterisks
    mark footnotes). We handle both.
    """
    if cell is None:
        return None
    if isinstance(cell, datetime):
        return cell.date()
    if isinstance(cell, date):
        return cell
    s = str(cell).strip()
    # Strip trailing non-digit / non-slash / non-dash characters (footnote markers)
    s = re.sub(r"[^0-9/\-]+$", "", s)
    if not s:
        return None
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%m-%d-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _to_decimal(v) -> Optional[Decimal]:
    if v is None:
        return None
    if isinstance(v, Decimal):
        return v
    if isinstance(v, (int, float)):
        return Decimal(str(v))
    s = str(v).strip()
    if not s:
        return None
    # Strip common currency / thousands formatting
    s = s.replace("$", "").replace(",", "").replace(" ", "")
    try:
        return Decimal(s)
    except (InvalidOperation, ValueError):
        return None


def _clean_str(v) -> Optional[str]:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


async def load_apg_base_rates_v2_from_bytes(
    session: AsyncSession,
    file_bytes: bytes,
    *,
    filename: str,
) -> dict:
    """Parse the workbook and replace all source='dtc' APG base rates.

    Returns a summary dict with:
      rows_deleted    — prior source='dtc' rows removed
      rows_inserted   — new rows written
      peer_groups     — distinct peer groups represented
      effective_dates — distinct effective dates represented
    """
    try:
        wb = openpyxl.load_workbook(
            io.BytesIO(file_bytes), read_only=True, data_only=True,
        )
    except Exception as e:
        raise ValueError(f"Could not open workbook: {e}") from e

    try:
        if SHEET_NAME not in wb.sheetnames:
            raise ValueError(
                f"Sheet {SHEET_NAME!r} not found in workbook. "
                f"Got sheets: {wb.sheetnames}"
            )
        ws = wb[SHEET_NAME]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()

    if len(rows) < DATA_START_ROW + 1:
        raise ValueError(
            f"{SHEET_NAME!r} sheet too short "
            f"(expected header at row {HEADER_ROW + 1}, data from row {DATA_START_ROW + 1})."
        )

    # Validate the header row
    header = rows[HEADER_ROW]
    if not header or _clean_str(header[0]) != "Peer Group":
        raise ValueError(
            f"{SHEET_NAME!r}: expected 'Peer Group' in row {HEADER_ROW + 1} "
            f"column A; got {header[0]!r}."
        )

    # Parse effective-date header columns
    date_columns: list[tuple[int, date]] = []
    for col_idx in range(FIRST_DATE_COL, len(header)):
        d = _parse_date_header(header[col_idx])
        if d is not None:
            date_columns.append((col_idx, d))
    if not date_columns:
        raise ValueError(
            f"{SHEET_NAME!r}: no effective-date columns detected in header."
        )
    log.info(
        "APG Base Rates v2: %d effective-date columns (oldest=%s, newest=%s)",
        len(date_columns),
        date_columns[0][1].isoformat(),
        date_columns[-1][1].isoformat(),
    )

    # Replace all source='dtc' rows.  (Hospital base rates untouched.)
    del_res = await session.execute(delete(ApgBaseRate).where(ApgBaseRate.source == "dtc"))
    rows_deleted = del_res.rowcount or 0

    new_rows: list[ApgBaseRate] = []
    peer_groups_seen: set[str] = set()
    effective_dates_seen: set[date] = set()

    for r_idx in range(DATA_START_ROW, len(rows)):
        row = rows[r_idx]
        if not row:
            continue
        peer_group = _clean_str(row[0])
        if not peer_group:
            # First blank row after the data block = stop.
            break
        if peer_group.startswith(_FOOTNOTE_PREFIX):
            # Hit the footnotes area.
            break

        base_rate_code = _clean_str(row[1]) if len(row) > 1 else None
        blend_rate_code = _clean_str(row[2]) if len(row) > 2 else None
        capital_rate_code = _clean_str(row[3]) if len(row) > 3 else None
        region = _clean_str(row[4]) if len(row) > 4 else None

        if not region:
            log.warning(
                "APG Base Rates v2: row %d (%r) has no region; skipping.",
                r_idx + 1, peer_group,
            )
            continue

        peer_groups_seen.add(peer_group)

        for col_idx, eff_date in date_columns:
            if col_idx >= len(row):
                continue
            rate = _to_decimal(row[col_idx])
            if rate is None or rate <= 0:
                # Empty / zero cells mean "rate not in effect this period."
                continue
            effective_dates_seen.add(eff_date)
            new_rows.append(ApgBaseRate(
                source="dtc",
                peer_group=peer_group,
                base_rate_code=base_rate_code,
                blend_rate_code=blend_rate_code,
                capital_rate_code=capital_rate_code,
                region=region,
                cheat_flag=None,
                cure_code=None,
                effective_date=eff_date,
                rate=rate,
            ))

    if not new_rows:
        raise ValueError(
            f"{SHEET_NAME!r}: parsed 0 usable rate rows. "
            f"File layout may have changed."
        )

    # Batch insert (typical payload is ~200 rows, but keep the batching
    # helper so the code tolerates a much larger file down the road).
    CHUNK = 2000
    for start in range(0, len(new_rows), CHUNK):
        session.add_all(new_rows[start:start + CHUNK])
        await session.flush()

    log.info(
        "APG Base Rates v2: deleted %d old, inserted %d new "
        "(%d peer groups × %d dates from %s → %s) from %r",
        rows_deleted, len(new_rows),
        len(peer_groups_seen), len(effective_dates_seen),
        min(effective_dates_seen).isoformat() if effective_dates_seen else "-",
        max(effective_dates_seen).isoformat() if effective_dates_seen else "-",
        filename,
    )

    return {
        "rows_deleted": rows_deleted,
        "rows_inserted": len(new_rows),
        "peer_groups": len(peer_groups_seen),
        "effective_dates": len(effective_dates_seen),
        "most_recent_effective_date": (
            max(effective_dates_seen).isoformat() if effective_dates_seen else None
        ),
    }
